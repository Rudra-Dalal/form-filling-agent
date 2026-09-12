import docx
from pathlib import Path

def parse_docx(file_path: str | Path) -> str:
    """Extract raw text from a DOCX file including paragraphs and tables."""
    doc = docx.Document(str(file_path))
    lines: list[str] = []

    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            lines.append(text)

    for table in doc.tables:
        for row in table.rows:
            row_vals = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_vals:
                lines.append(" : ".join(row_vals))

    return "\n".join(lines)
